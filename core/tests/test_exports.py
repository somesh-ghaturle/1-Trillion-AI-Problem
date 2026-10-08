import io
import json

import pandas as pd
from django.test import TestCase
from openpyxl import load_workbook

from core.models import DataSource, GovernanceMetric, ReconciliationRun, TrustScore


class ExportTest(TestCase):
    def setUp(self):
        source = DataSource.objects.create(name='=cmd|calc')
        TrustScore.objects.create(source=source, overall_score=80, trust_level='high', issues=['late data'])
        metric = GovernanceMetric.objects.create(name='rev', display_name='Revenue', description='-1 if refund',
                                                 data_type='numeric', formula='SUM(x)')
        ReconciliationRun.objects.create(governance_metric=metric, status='divergent', divergences=[
            {'source_a': 'A', 'source_b': '@B', 'severity': 'high', 'detail': 'differs'}])

    def get(self, dataset, fmt):
        resp = self.client.get(f'/export/{dataset}.{fmt}')
        self.assertEqual(resp.status_code, 200)
        self.assertIn(f'filename="{dataset}.{fmt}"', resp['Content-Disposition'])
        return resp.content

    def test_csv_neutralizes_formulas(self):
        df = pd.read_csv(io.BytesIO(self.get('trust-scores', 'csv')))
        self.assertEqual(df.loc[0, 'source_name'], "'=cmd|calc")
        self.assertEqual(json.loads(df.loc[0, 'issues']), ['late data'])

    def test_excel_cells_are_text_not_formulas(self):
        ws = load_workbook(io.BytesIO(self.get('governance-metrics', 'xlsx'))).active
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(row['description'], "'-1 if refund")
        self.assertEqual(ws.cell(2, header.index('description') + 1).data_type, 's')

    def test_parquet_keeps_raw_values_and_types(self):
        df = pd.read_parquet(io.BytesIO(self.get('trust-scores', 'parquet')))
        self.assertEqual(df.loc[0, 'source_name'], '=cmd|calc')
        self.assertEqual(df['overall_score'].dtype.kind, 'f')
        self.assertTrue(pd.api.types.is_datetime64_any_dtype(df['calculated_at']))

    def test_every_dataset_and_format(self):
        for dataset in ['trust-scores', 'validations', 'governance-metrics', 'semantic-definitions', 'reconciliation']:
            for fmt in ['csv', 'xlsx', 'parquet']:
                with self.subTest(dataset=dataset, fmt=fmt):
                    self.get(dataset, fmt)

    def test_unknown_dataset_or_format_404(self):
        self.assertEqual(self.client.get('/export/users.csv').status_code, 404)
        self.assertEqual(self.client.get('/export/trust-scores.pdf').status_code, 404)
